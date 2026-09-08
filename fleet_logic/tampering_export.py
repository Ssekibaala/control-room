"""
On-demand, filtered Tampering Risk Report export - the flexible
counterpart to tamper_engine.build_workbook(), which always produces
the full, fleet-wide, everything-included workbook.

This module never re-runs tampering analysis itself. It takes the SAME
per-case rows the dashboard's Confirmed/Unconfirmed Cases tables
already render (control_room._tamper_row()'s output, read straight out
of a dashboard-data payload) and just narrows the list down - by date
range, by client, by specific plates, or to confirmed-only - before
writing it out. That's deliberate: filtering already-classified rows
can never disagree with what the dashboard itself is showing for the
same cases, and it means a custom export never has to touch the
mailed-report CSVs or re-run tamper_engine.analyse() at request time.
"""

import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

NAVY = "1F3864"
WHITE = "FFFFFF"
_DATE_FMT = "%d/%m/%Y"

_COLUMNS = [
    {"key": "client", "label": "Client"},
    {"key": "plate", "label": "Plate"},
    {"key": "vehicle", "label": "Vehicle"},
    {"key": "fleetNumber", "label": "Fleet Number"},
    {"key": "severity", "label": "Severity"},
    {"key": "arrivalDate", "label": "Arrival Date"},
    {"key": "arrivalTime", "label": "Arrival Time"},
    {"key": "atLocation", "label": "At Location"},
    {"key": "nextDate", "label": "Next Date"},
    {"key": "nextTime", "label": "Next Time"},
    {"key": "fromLocation", "label": "From Location"},
    {"key": "distanceKm", "label": "Distance (km)"},
    {"key": "gapDuration", "label": "Gap Duration"},
    {"key": "impliedSpeed", "label": "Implied Speed (km/h)"},
    {"key": "powerEvent", "label": "Power Event"},
]


def parse_arrival_date(value):
    """arrivalDate is 'dd/mm/yyyy' (tamper_engine.DATE_TIME_FORMAT's date
    half) - returns None rather than raising on anything else, so a
    row with a malformed/missing date is excluded by a date filter
    instead of crashing the export."""
    try:
        return datetime.strptime((value or "").strip(), _DATE_FMT).date()
    except ValueError:
        return None


def filter_cases(cases, plate_owner, allowed_clients=None, client_name=None,
                  date_from=None, date_to=None, plates=None):
    """
    cases: a list of control_room._tamper_row()-shaped dicts (i.e.
    DATA.tamperConfirmed or DATA.tamperUnconfirmed straight from a
    dashboard-data payload).
    plate_owner: {plate: client}, see app.py's _plate_client_map() -
    tamper cases carry no client of their own, same reason that
    function exists for feedback comments.
    allowed_clients: the session's own visibility scope (None = every
    client, matching _visible_clients_for_session()) - always enforced
    first, regardless of what client_name asks for, so this can never
    be used to see a client's tampering cases the session isn't
    entitled to.
    client_name: further narrows to exactly one client, must itself be
    inside allowed_clients if both are given.
    date_from/date_to: inclusive, filtered against arrivalDate.
    plates: an iterable of plate strings (normalized/uppercased before
    comparing) - only these vehicles are included; falsy means every
    vehicle in scope.
    """
    plates_set = {p.strip().upper() for p in (plates or []) if p and p.strip()}
    out = []
    for c in cases:
        plate = (c.get("plate") or "").strip()
        owner = plate_owner.get(plate, "")
        if allowed_clients is not None and owner not in allowed_clients:
            continue
        if client_name and owner != client_name:
            continue
        if plates_set and plate.upper() not in plates_set:
            continue
        if date_from or date_to:
            d = parse_arrival_date(c.get("arrivalDate"))
            if d is None:
                continue
            if date_from and d < date_from:
                continue
            if date_to and d > date_to:
                continue
        out.append({**c, "client": owner})
    return out


def _write_sheet(wb, title, rows):
    ws = wb.create_sheet(title=title[:31])
    labels = [c["label"] for c in _COLUMNS]
    ws.append(labels)
    for col in range(1, len(_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if not rows:
        ws.append(["No cases match the selected filters."])
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(_COLUMNS))
    else:
        for r in rows:
            ws.append([r.get(c["key"], "") for c in _COLUMNS])
    ws.freeze_panes = "A2"
    if ws.max_row > 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(_COLUMNS))}{ws.max_row}"
    for i, c in enumerate(_COLUMNS, start=1):
        col = get_column_letter(i)
        cell_lens = [len(str(ws.cell(row=r, column=i).value or "")) for r in range(2, ws.max_row + 1)]
        ws.column_dimensions[col].width = min(max([len(c["label"])] + cell_lens) + 2, 46)
    return ws


def build_workbook(confirmed_cases, unconfirmed_cases, confirmed_only=False,
                    client_name=None, report_date=None):
    """
    confirmed_cases/unconfirmed_cases: already filtered (see
    filter_cases() above) - this function only lays them out, it makes
    no filtering decisions of its own.
    confirmed_only: when true, the Unconfirmed Cases sheet is omitted
    entirely rather than written out empty.
    """
    report_date = report_date or datetime.now()
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    title_ws = wb.create_sheet(title="Summary")
    title_ws["A1"] = "Tampering Risk Report" + (f" - {client_name}" if client_name else "")
    title_ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    title_ws["A2"] = report_date.strftime("%d %B %Y, %H:%M")
    title_ws["A2"].font = Font(italic=True, color="666666")
    title_ws["A4"] = "Confirmed cases"
    title_ws["B4"] = len(confirmed_cases)
    if not confirmed_only:
        title_ws["A5"] = "Unconfirmed cases"
        title_ws["B5"] = len(unconfirmed_cases)
    for col, w in (("A", 24), ("B", 12)):
        title_ws.column_dimensions[col].width = w

    _write_sheet(wb, "Confirmed Cases", confirmed_cases)
    if not confirmed_only:
        _write_sheet(wb, "Unconfirmed Cases", unconfirmed_cases)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
